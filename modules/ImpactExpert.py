import io
import re
from numbers import Number

import pandas as pd
import streamlit as st

from dataManager import load_all_data
from dbConfig import get_db_connection

BOBOT_SCHEME = [
    (0, 200_000_000, 0.0, 5.0),
    (200_000_000, 1_000_000_000, 5.1, 10.0),
    (1_000_000_000, 5_000_000_000, 10.1, 15.0),
]
BOBOT_MAKSIMAL = BOBOT_SCHEME[-1][3]

KOLOM_DIBUTUHKAN = {
    "nik": ["NIK"],
    "expert": ["NAMA"],
    "assignment": ["ASSIGNMENT"],
    "financial_eksisting": ["FINANCIAL", "EKSISTING"],
    "potensi_revenue": ["POT", "REVENUE"],
}

LABEL_KOLOM = {
    "nik": "NIK",
    "expert": "Nama Expert",
    "assignment": "Assignment/Penugasan",
    "financial_eksisting": "Revenue Eksisting (per Assignment)",
    "potensi_revenue": "Potensi Revenue Tambahan (per Assignment)",
    "total_financial_eksisting": "Total Revenue Eksisting (Semua Assignment)",
    "hasil_impact_expert": "Hasil Impact Expert (%)",
    "jumlah_assignment": "Jumlah Assignment",
}

KUNCI_ASSIGNMENT = ["nik", "expert", "assignment"]
KUNCI_EXPERT = ["nik", "expert"]

RUPIAH_PREFIX_RE = re.compile(r"(?i)^rp\.?")
NAMA_TABEL_DB = "expert_impact"
KOLOM_DB_KE_INTERNAL = {
    "customer": "assignment",
    "pot_revenue_tambahan": "potensi_revenue",
}

LABEL_TOTAL = {
    "TOTAL",
    "GRAND TOTAL",
    "TOTAL REVENUE",
    "TOTAL REVENUE EKSISTING",
    "JUMLAH TOTAL",
}

KOLOM_HASIL = [
    "blok_assignment",
    "nik",
    "expert",
    "assignment",
    "financial_eksisting",
    "potensi_revenue",
    "quarter",
]


def hitung_bobot_impact(nilai):
    """Hitung Bobot Impact Expert (%) dari total Financial Eksisting."""
    if pd.isna(nilai) or nilai <= 0:
        return 0.0

    for batas_bawah, batas_atas, bobot_bawah, bobot_atas in BOBOT_SCHEME:
        if nilai <= batas_atas:
            proporsi = (nilai - batas_bawah) / (batas_atas - batas_bawah)
            return bobot_bawah + proporsi * (bobot_atas - bobot_bawah)

    return BOBOT_MAKSIMAL


def _is_blank(value):
    if value is None:
        return True
    if isinstance(value, str):
        return not value.strip()
    try:
        return bool(pd.isna(value))
    except Exception:
        return False


def _parse_rupiah(value):
    """
    Parse nominal Financial/Potensi dari Excel atau CSV.

    Semua bentuk berikut dihitung:
      396000000
      396000000.00
      "38.000.000.000"
      "Rp38.000.000.000"
      "Rp. 24.000.000.000"
      "Rp. 100.000.000,-"

    STRING rupiah tidak diabaikan, karena pada file Impact yang bersih sebagian
    nominal memang tersimpan sebagai text, misalnya "38.000.000.000".
    """
    if _is_blank(value):
        return 0.0

    if isinstance(value, Number) and not isinstance(value, bool):
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0

    text = str(value).strip()
    if not text or text == "-":
        return 0.0

    # Formula jangan dianggap sebagai nominal data.
    if text.startswith("="):
        return 0.0

    text = text.replace(" ", "")
    text = RUPIAH_PREFIX_RE.sub("", text)
    text = text.rstrip(",-")
    text = text.strip(".")

    if not text:
        return 0.0

    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif text.count(".") > 1:
        text = text.replace(".", "")
    elif "," in text:
        head, _, tail = text.rpartition(",")
        if head and len(tail) in (1, 2):
            text = f"{head.replace(',', '')}.{tail}"
        else:
            text = text.replace(",", "")
    elif text.count(".") == 1:
        head, tail = text.split(".", 1)
        if len(tail) not in (1, 2):
            text = head + tail

    try:
        return float(text)
    except ValueError:
        return 0.0


def _format_rupiah(x):
    return f"Rp {float(x):,.0f}".replace(",", ".")


def _normalize_header(value):
    if _is_blank(value):
        return ""
    return " ".join(str(value).upper().split())


def _clean_text(value):
    if _is_blank(value):
        return None
    cleaned = " ".join(str(value).split())
    return cleaned or None


def _clean_nik(value):
    cleaned = _clean_text(value)
    if cleaned is None:
        return None
    if cleaned.endswith(".0"):
        cleaned = cleaned[:-2]
    return cleaned


_MULTI_VALUE_SPLIT_RE = re.compile(r"[\r\n]+")


def _split_multi_values(value):
    """
    Pecah satu cell yang berisi beberapa nilai sekaligus (dipisah baris baru
    di dalam cell yang sama, mis. beberapa NIK/Nama expert untuk satu
    assignment yang sama). Nilai tunggal tanpa baris baru tetap menghasilkan
    list berisi satu elemen, sehingga logika downstream selalu bekerja
    dengan list.
    """
    if _is_blank(value):
        return []
    parts = _MULTI_VALUE_SPLIT_RE.split(str(value))
    return [p.strip() for p in parts if p.strip()]


def _split_nik_expert_pair(nik_value, expert_value, cleaner_nik, cleaner_expert):
    """
    Pecah pasangan cell NIK & Nama Expert yang mungkin berisi beberapa
    peserta sekaligus, lalu dipasangkan berdasarkan urutan baris di dalam
    cell. Jika jumlah baris NIK dan Nama tidak sama, sisanya dipasangkan
    dengan None supaya tidak ada data yang salah tertukar pasangannya.
    """
    nik_parts = [cleaner_nik(p) for p in _split_multi_values(nik_value)]
    expert_parts = [cleaner_expert(p) for p in _split_multi_values(expert_value)]

    if not nik_parts and not expert_parts:
        return [], []

    n = max(len(nik_parts), len(expert_parts))
    nik_parts += [None] * (n - len(nik_parts))
    expert_parts += [None] * (n - len(expert_parts))
    return nik_parts, expert_parts


def _find_column(headers, required_substrings):
    for idx, header in enumerate(headers):
        if all(sub in header for sub in required_substrings):
            return idx
    return None


def _locate_columns(headers):
    found, missing = {}, []
    for name, keywords in KOLOM_DIBUTUHKAN.items():
        idx = _find_column(headers, keywords)
        if idx is None:
            missing.append(name)
        else:
            found[name] = idx
    return found, missing


HEADER_SCAN_MAX_ROWS = 8


def _detect_header_rows(raw):
    """
    Deteksi baris awal header (bisa satu atau dua baris), di mana pun header
    itu berada dalam beberapa baris pertama file.

    Beberapa export Excel menaruh header persis di baris pertama, tapi ada
    juga yang menaruh 1-2 baris kosong/judul di atasnya dulu (mis. header
    baru muncul di baris ke-2 dan ke-3). Fungsi ini mencoba tiap kemungkinan
    baris awal sampai ditemukan kombinasi yang berhasil memetakan seluruh
    kolom yang dibutuhkan.

    Return: (headers, header_start_row, n_header_rows)
      - header_start_row: index 0-based dari raw, baris pertama header.
      - n_header_rows: 1 atau 2, jumlah baris yang membentuk header.
    """
    limit = min(HEADER_SCAN_MAX_ROWS, len(raw))
    best = None  # (missing_count, headers, start, n_rows)

    for start in range(limit):
        headers_1 = [_normalize_header(v) for v in raw.iloc[start]]
        _, missing_1 = _locate_columns(headers_1)
        if not missing_1:
            return headers_1, start, 1
        if best is None or len(missing_1) < best[0]:
            best = (len(missing_1), headers_1, start, 1)

        if start + 1 < len(raw):
            headers_2 = [
                _normalize_header(sub) or _normalize_header(top)
                for top, sub in zip(raw.iloc[start], raw.iloc[start + 1])
            ]
            _, missing_2 = _locate_columns(headers_2)
            if not missing_2:
                return headers_2, start, 2
            if best is None or len(missing_2) < best[0]:
                best = (len(missing_2), headers_2, start, 2)

    # Tidak ada kombinasi yang sempurna: kembalikan kandidat terbaik supaya
    # pesan error "kolom tidak ditemukan" tetap informatif.
    _, headers, start, n_rows = best
    return headers, start, n_rows


def _bangun_hasil_dari_baris(baris, nama_file, quarter_upload):
    """
    Ubah baris mentah (satu baris Excel/CSV) menjadi tabel hasil per expert.

    Logika ini dipakai bersama oleh parser XLSX dan CSV supaya tidak ada lagi
    versi yang tertinggal saat salah satunya diubah.
    """
    if baris is None or baris.empty:
        return None

    baris = baris.copy()

    # Baris dengan Assignment terisi menandai awal satu blok assignment baru.
    assignment_baru = baris["assignment_baris"].notna()
    if len(assignment_baru) and not assignment_baru.iloc[0]:
        assignment_baru.iloc[0] = True

    baris["blok_assignment_num"] = assignment_baru.cumsum()
    baris["assignment"] = baris["assignment_baris"].ffill()

    # Buang baris yang tidak memiliki assignment sama sekali.
    baris = baris[baris["assignment"].notna()].copy()
    if baris.empty:
        return None

    # Total Financial/Potensi per SATU kejadian assignment.
    total_per_blok = (
        baris.groupby("blok_assignment_num", as_index=False, dropna=False)
        .agg(
            assignment=("assignment", "first"),
            financial_eksisting=("financial_eksisting", "sum"),
            potensi_revenue=("potensi_revenue", "sum"),
        )
    )

    ada_identitas = baris["nik_baris"].apply(bool) | baris["expert_baris"].apply(bool)
    peserta = baris.loc[
        ada_identitas,
        ["blok_assignment_num", "nik_baris", "expert_baris"],
    ].copy()

    if peserta.empty:
        return None

    peserta = peserta.explode(
        ["nik_baris", "expert_baris"], ignore_index=True
    )

    peserta["nik"] = pd.to_numeric(peserta["nik_baris"], errors="coerce")
    peserta = peserta.dropna(subset=["nik", "expert_baris"])
    peserta = peserta.drop_duplicates(subset=["blok_assignment_num", "nik"])
    peserta = peserta.rename(columns={"expert_baris": "expert"})[
        ["blok_assignment_num", "nik", "expert"]
    ]

    if peserta.empty:
        return None

    # Fan-out: satu assignment menjadi satu nilai penuh bagi setiap expert.
    hasil = peserta.merge(
        total_per_blok,
        on="blok_assignment_num",
        how="left",
    )
    hasil["blok_assignment"] = (
        str(nama_file) + "::" + hasil["blok_assignment_num"].astype(str)
    )
    hasil["quarter"] = quarter_upload

    return hasil[KOLOM_HASIL]


def _read_xlsx_as_raw(uploaded_file):
    """
    Membaca sheet pertama dan mempertahankan apakah cell berisi numeric, text,
    atau formula. Ini penting agar kita dapat membedakan formula total dari
    nominal data biasa.
    """
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("Package openpyxl diperlukan untuk membaca XLSX.") from exc

    workbook = load_workbook(
        filename=io.BytesIO(uploaded_file.getvalue()),
        data_only=False,
    )
    worksheet = workbook[workbook.sheetnames[0]]

    values = []
    for row in worksheet.iter_rows(values_only=False):
        current = []
        for cell in row:
            current.append(cell.value)
        values.append(current)

    return pd.DataFrame(values, dtype=object), workbook, worksheet


def _xlsx_row_is_summary_or_formula(worksheet, excel_row_number, financial_col_idx):
    """
    Identifikasi baris total/summary.

    Pembacaan data dihentikan jika Financial Eksisting pada baris tersebut
    adalah formula (contoh: =SUM(...)), atau terdapat label TOTAL/GRAND TOTAL.
    Dengan cara ini formula total Excel tidak dijumlahkan lagi sebagai data.
    """
    financial_cell = worksheet.cell(
        row=excel_row_number,
        column=financial_col_idx + 1,
    )
    if financial_cell.data_type == "f":
        return True
    if isinstance(financial_cell.value, str) and financial_cell.value.strip().startswith("="):
        return True

    for col_idx in range(1, min(worksheet.max_column, 5) + 1):
        value = worksheet.cell(excel_row_number, col_idx).value
        if isinstance(value, str):
            text = " ".join(value.upper().split())
            if text in LABEL_TOTAL:
                return True
            if "GRAND TOTAL" in text or text.startswith("TOTAL REVENUE"):
                return True

    return False


def _parse_satu_file_xlsx(uploaded_file, quarter_upload):
    raw, _workbook, worksheet = _read_xlsx_as_raw(uploaded_file)

    if len(raw) < 3:
        st.warning(f"File {uploaded_file.name} tidak punya cukup baris data.")
        return None

    headers, header_start, n_header_rows = _detect_header_rows(raw)
    columns, missing = _locate_columns(headers)
    if missing:
        st.error(
            f"Kolom berikut tidak ditemukan di {uploaded_file.name}: "
            f"{', '.join(LABEL_KOLOM.get(m, m) for m in missing)}. "
            f"Header yang terbaca: {headers}"
        )
        return None

    # Data hanya dari bawah header sampai sebelum formula/summary.
    start_excel_row = header_start + n_header_rows + 1  # 1-based excel row
    data_records = []

    for excel_row_number in range(start_excel_row, worksheet.max_row + 1):
        if _xlsx_row_is_summary_or_formula(
            worksheet,
            excel_row_number,
            columns["financial_eksisting"],
        ):
            break

        row_idx = excel_row_number - 1
        if row_idx >= len(raw):
            break

        raw_row = raw.iloc[row_idx]

        # Baris kosong total: tetap dicatat agar blok assignment tidak terputus.
        if all(_is_blank(raw_row.iloc[i]) for i in range(len(raw_row))):
            data_records.append(
                {
                    "nik_baris": [],
                    "expert_baris": [],
                    "assignment_baris": None,
                    "financial_eksisting": 0.0,
                    "potensi_revenue": 0.0,
                }
            )
            continue

        nik_parts, expert_parts = _split_nik_expert_pair(
            raw_row.iloc[columns["nik"]],
            raw_row.iloc[columns["expert"]],
            _clean_nik,
            _clean_text,
        )

        data_records.append(
            {
                "nik_baris": nik_parts,
                "expert_baris": expert_parts,
                "assignment_baris": _clean_text(raw_row.iloc[columns["assignment"]]),
                "financial_eksisting": _parse_rupiah(
                    raw_row.iloc[columns["financial_eksisting"]]
                ),
                "potensi_revenue": _parse_rupiah(
                    raw_row.iloc[columns["potensi_revenue"]]
                ),
            }
        )

    return _bangun_hasil_dari_baris(
        pd.DataFrame(data_records),
        uploaded_file.name,
        quarter_upload,
    )


def _parse_satu_file_csv(uploaded_file, quarter_upload):
    raw = pd.read_csv(uploaded_file, header=None, dtype=object)

    if len(raw) < 3:
        st.warning(f"File {uploaded_file.name} tidak punya cukup baris data.")
        return None

    headers, header_start, n_header_rows = _detect_header_rows(raw)
    columns, missing = _locate_columns(headers)
    if missing:
        st.error(
            f"Kolom berikut tidak ditemukan di {uploaded_file.name}: "
            f"{', '.join(LABEL_KOLOM.get(m, m) for m in missing)}."
        )
        return None

    data = raw.iloc[header_start + n_header_rows:].reset_index(drop=True)

    records = []
    for _, row in data.iterrows():
        financial_raw = row.iloc[columns["financial_eksisting"]]
        text_financial = (
            " ".join(str(financial_raw).upper().split())
            if not _is_blank(financial_raw)
            else ""
        )
        if text_financial.startswith("="):
            break
        if text_financial in LABEL_TOTAL:
            break

        nik_parts, expert_parts = _split_nik_expert_pair(
            row.iloc[columns["nik"]],
            row.iloc[columns["expert"]],
            _clean_nik,
            _clean_text,
        )

        records.append(
            {
                "nik_baris": nik_parts,
                "expert_baris": expert_parts,
                "assignment_baris": _clean_text(row.iloc[columns["assignment"]]),
                "financial_eksisting": _parse_rupiah(financial_raw),
                "potensi_revenue": _parse_rupiah(
                    row.iloc[columns["potensi_revenue"]]
                ),
            }
        )

    return _bangun_hasil_dari_baris(
        pd.DataFrame(records),
        uploaded_file.name,
        quarter_upload,
    )


def _parse_satu_file(uploaded_file, quarter_upload):
    name = uploaded_file.name.lower()
    if name.endswith(".xlsx"):
        return _parse_satu_file_xlsx(uploaded_file, quarter_upload)
    if name.endswith(".xls"):
        st.error("Format .xls lama tidak didukung. Simpan file sebagai .xlsx atau .csv.")
        return None
    return _parse_satu_file_csv(uploaded_file, quarter_upload)


def read_and_merge(files, quarter_upload):
    all_rows = []

    for uploaded_file in files:
        try:
            df = _parse_satu_file(uploaded_file, quarter_upload)
        except Exception as exc:
            st.error(f"Gagal membaca file {uploaded_file.name}: {exc}")
            continue

        if df is not None and not df.empty:
            all_rows.append(df)

    if not all_rows:
        return pd.DataFrame()

    combined = pd.concat(all_rows, ignore_index=True)

    # Pertahankan nama expert yang paling awal untuk NIK yang sama.
    combined["expert"] = combined.groupby("nik")["expert"].transform("first")
    return combined


def build_raw_table(df):
    """
    SATU BARIS = SATU KEJADIAN ASSIGNMENT.

    Ini adalah sumber angka untuk Total Revenue Eksisting dashboard.
    Revenue dari beberapa baris di dalam satu merged assignment dijumlahkan
    sekali per assignment, lalu TIDAK dikalikan dengan jumlah expert.
    """
    if df.empty:
        return pd.DataFrame(
            columns=[
                "assignment",
                "financial_eksisting",
                "potensi_revenue",
            ]
        )

    if "blok_assignment" in df.columns:
        # Setelah fan-out, assignment yang sama muncul berkali-kali per expert.
        # Ambil hanya satu salinan per kejadian assignment.
        unique = df.drop_duplicates(subset=["blok_assignment"]).copy()
    else:
        # Fallback data lama dari DB.
        unique = df.drop_duplicates(
            subset=["assignment", "financial_eksisting", "potensi_revenue"]
        ).copy()

    return unique[
        ["assignment", "financial_eksisting", "potensi_revenue"]
    ].reset_index(drop=True)


def build_assignment_table(df):
    """Satu baris per (NIK, Expert, Assignment)."""
    if df.empty:
        return pd.DataFrame(
            columns=KUNCI_ASSIGNMENT
            + ["financial_eksisting", "potensi_revenue"]
        )

    return (
        df.groupby(KUNCI_ASSIGNMENT, as_index=False, dropna=False)
        .agg(
            financial_eksisting=("financial_eksisting", "sum"),
            potensi_revenue=("potensi_revenue", "sum"),
        )
    )


def compute_impact_scores(assignment_df):
    """Jumlahkan seluruh assignment per expert lalu hitung Impact Expert."""
    if assignment_df.empty:
        return pd.DataFrame(
            columns=KUNCI_EXPERT
            + [
                "financial_eksisting",
                "potensi_revenue",
                "jumlah_assignment",
                "hasil_impact_expert",
            ]
        )

    summary = (
        assignment_df.groupby(KUNCI_EXPERT, as_index=False, dropna=False)
        .agg(
            financial_eksisting=("financial_eksisting", "sum"),
            potensi_revenue=("potensi_revenue", "sum"),
            jumlah_assignment=("assignment", "nunique"),
        )
    )

    summary["hasil_impact_expert"] = (
        summary["financial_eksisting"]
        .apply(hitung_bobot_impact)
        .round(3)
    )

    return summary.sort_values(
        "financial_eksisting", ascending=False
    ).reset_index(drop=True)


def build_impact_detail_table(assignment_df, summary):
    detail = assignment_df.merge(
        summary[
            KUNCI_EXPERT
            + ["financial_eksisting", "hasil_impact_expert"]
        ].rename(
            columns={
                "financial_eksisting": "total_financial_eksisting"
            }
        ),
        on=KUNCI_EXPERT,
        how="left",
    )

    urutan_expert = [
        (row["nik"], row["expert"])
        for _, row in summary.iterrows()
    ]

    detail["_urutan_expert"] = pd.Categorical(
        list(zip(detail["nik"], detail["expert"])),
        categories=urutan_expert,
        ordered=True,
    )

    detail = (
        detail.sort_values(
            ["_urutan_expert", "financial_eksisting"],
            ascending=[True, False],
        )
        .drop(columns="_urutan_expert")
        .reset_index(drop=True)
    )

    return detail[
        [
            "nik",
            "expert",
            "assignment",
            "financial_eksisting",
            "potensi_revenue",
            "total_financial_eksisting",
            "hasil_impact_expert",
        ]
    ]


def save_scores_to_db(supabase, summary, quarter):
    for _, row in summary.iterrows():
        supabase.table("calculated").upsert(
            {
                "nik": int(row["nik"]),
                "expert": row["expert"],
                "quarter": quarter,
                "impact_expert": float(row["hasil_impact_expert"]),
            },
            on_conflict="nik,expert,quarter",
        ).execute()

    return len(summary)


def _render_raw_table(raw_df):
    st.subheader("🗂️ Raw Data (Assignment Unik)")
    st.caption(
        "Satu baris = satu kejadian assignment. Semua nominal Financial Eksisting "
        "di dalam blok assignment dijumlahkan sekali. Nominal tidak digandakan "
        "berdasarkan jumlah expert."
    )

    raw_display = raw_df.copy()
    for col in ["financial_eksisting", "potensi_revenue"]:
        if col in raw_display.columns:
            raw_display[col] = raw_display[col].apply(_format_rupiah)

    st.dataframe(
        raw_display.rename(columns=LABEL_KOLOM),
        use_container_width=True,
        hide_index=True,
    )


def _render_detail_table(assignment_df, summary):
    st.subheader("📋 Rincian Assignment per Expert")
    st.caption(
        "Setiap baris = satu assignment yang diikuti expert. Nilai assignment "
        "merupakan total seluruh baris Financial Eksisting di blok assignment."
    )

    detail_display = build_impact_detail_table(assignment_df, summary).copy()
    for col in [
        "financial_eksisting",
        "potensi_revenue",
        "total_financial_eksisting",
    ]:
        detail_display[col] = detail_display[col].apply(_format_rupiah)

    st.dataframe(
        detail_display.rename(columns=LABEL_KOLOM),
        use_container_width=True,
        hide_index=True,
    )


def _render_summary_table(summary):
    st.subheader("📊 Rekap Impact Expert per Expert")
    display_df = summary.copy()
    for col in ["financial_eksisting", "potensi_revenue"]:
        display_df[col] = display_df[col].apply(_format_rupiah)

    st.dataframe(
        display_df.rename(columns=LABEL_KOLOM),
        use_container_width=True,
        hide_index=True,
    )

    st.download_button(
        "Download rekap CSV",
        data=summary.rename(columns=LABEL_KOLOM)
        .to_csv(index=False)
        .encode("utf-8"),
        file_name="impact_expert_summary.csv",
        mime="text/csv",
    )


def _render_excel_download(assignment_df, summary, quarter):
    detail_download = build_impact_detail_table(
        assignment_df, summary
    ).rename(columns=LABEL_KOLOM)
    rekap_download = summary.rename(columns=LABEL_KOLOM)

    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        detail_download.to_excel(writer, sheet_name="Detail", index=False)
        rekap_download.to_excel(writer, sheet_name="Rekap", index=False)
    buffer.seek(0)

    st.download_button(
        label="Download Hasil Impact Expert (Excel)",
        data=buffer,
        file_name=f"Impact_Expert_{quarter}.xlsx",
        mime=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        key="btn_download_impact",
    )


def _load_data_dari_upload():
    quarter_upload = st.selectbox(
        "Quarter untuk data ini",
        ["Q1", "Q2", "Q3", "Q4"],
    )
    uploaded_files = st.file_uploader(
        "Upload data Impact Expert (format Excel/CSV)",
        accept_multiple_files=True,
        type=["xlsx", "csv"],
    )

    if uploaded_files:
        st.session_state["impact_combined_df"] = read_and_merge(
            uploaded_files,
            quarter_upload,
        )
        st.session_state["impact_quarter"] = quarter_upload

    combined_df = st.session_state.get(
        "impact_combined_df",
        pd.DataFrame(),
    )
    quarter = st.session_state.get(
        "impact_quarter",
        quarter_upload,
    )
    return combined_df, quarter


def _load_data_dari_database():
    combined_df = load_all_data(NAMA_TABEL_DB)
    if not combined_df.empty:
        combined_df = combined_df.rename(
            columns=KOLOM_DB_KE_INTERNAL
        )

    quarter = st.pills(
        "Pilih Quarter",
        ["Q1", "Q2", "Q3", "Q4"],
        selection_mode="single",
        default="Q1",
    )

    if "quarter" in combined_df.columns:
        combined_df = combined_df[
            combined_df["quarter"] == quarter
        ]

    return combined_df, quarter


def ImpactExpert():
    supabase = get_db_connection()

    st.title("💥 Impact Expert")
    st.caption(
        "Revenue Eksisting dihitung dari tabel Impact sebagai berikut: "
        "semua nominal Financial Eksisting di dalam satu blok assignment "
        "dijumlahkan sekali per assignment, termasuk nominal yang ditulis "
        "sebagai teks seperti '38.000.000.000'. Nominal tersebut baru "
        "di-fan-out ke expert untuk kebutuhan scoring, tetapi Total Revenue "
        "perusahaan tetap hanya menjumlahkan assignment unik."
    )

    mode = st.pills(
        "Data Resource",
        ["Upload file", "From Data Base"],
        selection_mode="single",
        default="From Data Base",
    )

    if mode == "Upload file":
        st.header("📁 Upload File")
        combined_df, quarter = _load_data_dari_upload()
    else:
        combined_df, quarter = _load_data_dari_database()

    if combined_df.empty:
        st.info(
            f"Tidak ada data untuk {quarter}. "
            "Pilih quarter lain atau upload file yang sesuai."
        )
        return

    # Fallback DB lama.
    if "assignment" not in combined_df.columns:
        combined_df["assignment"] = "-"

    missing_cols = [
        c
        for c in [
            "nik",
            "expert",
            "assignment",
            "financial_eksisting",
            "potensi_revenue",
        ]
        if c not in combined_df.columns
    ]
    if missing_cols:
        st.error(
            f"Kolom berikut tidak ditemukan: {missing_cols}."
        )
        return

    combined_df = combined_df.copy()

    # Pastikan nominal numeric.
    combined_df["financial_eksisting"] = combined_df[
        "financial_eksisting"
    ].apply(_parse_rupiah)
    combined_df["potensi_revenue"] = combined_df[
        "potensi_revenue"
    ].apply(_parse_rupiah)

    # Agregasi.
    raw_df = build_raw_table(combined_df)
    assignment_df = build_assignment_table(combined_df)
    summary = compute_impact_scores(assignment_df)

    total_revenue_eksisting = raw_df["financial_eksisting"].sum()
    total_revenue_fanout = assignment_df["financial_eksisting"].sum()

    col1, col2 = st.columns(2)
    col1.metric(
        "Total Expert",
        len(summary),
        border=True,
    )
    col2.metric(
        "Total Revenue Eksisting",
        _format_rupiah(total_revenue_eksisting),
        border=True,
    )

    with st.expander("🔎 Validasi Total Revenue", expanded=False):
        st.write(f"Jumlah assignment unik: **{len(raw_df)}**")
        st.write(
            "Total Revenue Eksisting dari assignment unik: "
            f"**{_format_rupiah(total_revenue_eksisting)}**"
        )
        st.write(
            "Total setelah fan-out ke expert (jangan dipakai sebagai total perusahaan): "
            f"**{_format_rupiah(total_revenue_fanout)}**"
        )
        st.caption(
            "Angka fan-out bisa lebih besar karena satu assignment dapat diikuti "
            "lebih dari satu expert."
        )

    _render_raw_table(raw_df)
    _render_detail_table(assignment_df, summary)
    _render_summary_table(summary)
    _render_excel_download(assignment_df, summary, quarter)

    if st.button(
        "💾 Simpan Impact Expert Score ke Database",
        key="btn_save_impact",
    ):
        try:
            count = save_scores_to_db(supabase, summary, quarter)
            st.session_state[
                "impact_save_message"
            ] = f"✅ {count} data disimpan/diupdate untuk {quarter}"
            st.rerun()
        except Exception as exc:
            st.error(f"❌ Error: {exc}")

    if "impact_save_message" in st.session_state:
        st.success(st.session_state.pop("impact_save_message"))


if __name__ == "__main__":
    ImpactExpert()